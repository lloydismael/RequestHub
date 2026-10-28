"""Regression coverage for independent SQR remarks and the manual hold workflow."""

import io
import json
from datetime import timedelta
from html.parser import HTMLParser
from unittest.mock import patch

from django.contrib.messages.storage.fallback import FallbackStorage
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.test import RequestFactory, TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from hub.models import Notification, SqrSubmission, SqrSubmissionChange, SqrSubmissionHistory
from hub.test_multiple_sqrs import MultipleSqrFixtures
from hub.views import SqrEngineerUpdateView, SqrImportView


class TrackerRows(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []
        self.row = None
        self.cell = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = {"attrs": dict(attrs), "cells": []}
        elif tag in ("th", "td") and self.row is not None:
            self.cell = {"tag": tag, "attrs": dict(attrs), "text": ""}
            self.row["cells"].append(self.cell)

    def handle_data(self, data):
        if self.cell is not None:
            self.cell["text"] += data

    def handle_endtag(self, tag):
        if tag in ("th", "td"):
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None


@override_settings(ACS_EMAIL_CONNECTION_STRING="", ACS_EMAIL_SENDER="")
class SqrOnHoldTests(MultipleSqrFixtures, TestCase):
    def setUp(self):
        super().setUp()
        self.paused_engineer = User.objects.create_user(
            username="hold_sqr_engineer", role=User.Roles.ON_HOLD,
        )
        self.submission = self.create_sqr()
        self.submission.remarks = "SSE filing remarks <independent>"
        self.submission.review_notes = "Existing reviewer notes"
        self.submission.save()
        self.client.force_login(self.pm)

    def update_status(self, value, *, submission=None, **extra):
        return self.client.post(
            reverse("hub:sqr-inline-update", args=[(submission or self.submission).pk]),
            data=json.dumps({"field": "status", "value": value, **extra}),
            content_type="application/json",
        )

    def snapshot(self):
        return SqrSubmission.objects.filter(pk=self.submission.pk).values().get()

    def assert_rejected_without_mutation(self, response, before, counts, status=400):
        self.assertEqual(response.status_code, status)
        self.assertFalse(response.json()["ok"])
        self.assertEqual(self.snapshot(), before)
        self.assertEqual(
            (SqrSubmissionChange.objects.count(), SqrSubmissionHistory.objects.count(), Notification.objects.count()),
            counts,
        )

    def mutation_counts(self):
        return (
            SqrSubmissionChange.objects.count(), SqrSubmissionHistory.objects.count(), Notification.objects.count(),
        )

    def test_model_has_distinct_optional_remarks_and_hold_reason(self):
        self.assertEqual(SqrSubmission.Status.ON_HOLD, "on_hold")
        self.assertEqual(SqrSubmission.Status.ON_HOLD.label, "On Hold")
        self.assertTrue(SqrSubmission._meta.get_field("remarks").blank)
        self.assertTrue(SqrSubmission._meta.get_field("hold_reason").blank)
        self.assertEqual(SqrSubmission._meta.get_field("hold_reason").get_internal_type(), "TextField")
        self.assertEqual(self.submission.hold_reason, "")

    def test_pm_and_admin_can_hold(self):
        for user in (self.pm, self.admin):
            with self.subTest(role=user.role):
                sqr = self.create_sqr()
                self.client.force_login(user)
                response = self.update_status("on_hold", submission=sqr, hold_reason="Awaiting customer scope")
                self.assertEqual(response.status_code, 200, response.content)
                self.assertTrue(response.json()["ok"])
                self.assertEqual(response.json()["hold_reason"], "Awaiting customer scope")
                sqr.refresh_from_db()
                self.assertEqual(sqr.status, "on_hold")
                self.assertEqual(sqr.hold_reason, "Awaiting customer scope")

    def test_other_roles_cannot_hold_or_resume(self):
        SqrSubmission.objects.filter(pk=self.submission.pk).update(
            status="on_hold", hold_reason="Existing hold",
        )
        for user in (self.engineer, self.other, self.paused_engineer, self.requestor):
            for target in ("on_hold", "submitted", "for_revision", "reviewed"):
                with self.subTest(role=user.role, target=target):
                    self.client.force_login(user)
                    before, counts = self.snapshot(), self.mutation_counts()
                    response = self.update_status(target, hold_reason="Unauthorized replacement")
                    self.assert_rejected_without_mutation(response, before, counts, status=403)

    def test_hold_requires_nonwhitespace_string_and_failure_is_atomic(self):
        invalid_payloads = [{}, *({"hold_reason": value} for value in ("", " \t\r\n ", None, 0, False, [], {}))]
        for payload in invalid_payloads:
            with self.subTest(payload=payload):
                before, counts = self.snapshot(), self.mutation_counts()
                response = self.update_status("on_hold", **payload)
                self.assert_rejected_without_mutation(response, before, counts)

    def test_old_reason_cannot_substitute_for_reason_on_next_hold(self):
        self.assertEqual(self.update_status("on_hold", hold_reason="Initial delay").status_code, 200)
        self.assertEqual(self.update_status("submitted").status_code, 200)
        before, counts = self.snapshot(), self.mutation_counts()
        self.assert_rejected_without_mutation(self.update_status("on_hold"), before, counts)
        response = self.update_status("on_hold", hold_reason="New delay")
        self.assertEqual(response.status_code, 200)
        self.submission.refresh_from_db()
        self.assertEqual(self.submission.hold_reason, "New delay")
        self.assertEqual(self.submission.remarks, "SSE filing remarks <independent>")

    def test_unknown_empty_and_nonstring_statuses_are_rejected_atomically(self):
        for status in ("", " ", "unknown", "approved", None, 0, [], {}):
            with self.subTest(status=status):
                before, counts = self.snapshot(), self.mutation_counts()
                response = self.update_status(status, hold_reason="Must not be saved")
                self.assert_rejected_without_mutation(response, before, counts)

    @patch("hub.views._send_sqr_approved_email")
    @patch("hub.views._send_sqr_for_revision_email")
    def test_pause_preserves_approval_metadata_and_has_no_notifications(self, revision_email, approval_email):
        reviewed_at = timezone.now() - timedelta(days=7)
        due_date = timezone.localdate() + timedelta(days=23)
        SqrSubmission.objects.filter(pk=self.submission.pk).update(
            status="reviewed", reviewed_at=reviewed_at, reviewed_by=self.pm,
            validity_due_date=due_date, assigned_pm=self.pm,
        )
        before = self.snapshot()
        notifications = Notification.objects.count()
        response = self.update_status("on_hold", hold_reason="Budget confirmation pending")
        self.assertEqual(response.status_code, 200, response.content)
        after = self.snapshot()
        for field in ("reviewed_at", "reviewed_by_id", "validity_due_date", "assigned_pm_id", "remarks", "review_notes"):
            self.assertEqual(after[field], before[field], field)
        revision_email.assert_not_called()
        approval_email.assert_not_called()
        self.assertEqual(Notification.objects.count(), notifications)
        self.assertFalse(response.json().get("approval_email_draft_url"))
        self.assertFalse(response.json().get("approval_email_eml_url"))

    def test_unapproved_pause_does_not_create_approval_or_revision_metadata(self):
        before = self.snapshot()
        response = self.update_status("on_hold", hold_reason="Awaiting scope")
        self.assertEqual(response.status_code, 200)
        after = self.snapshot()
        for field in ("reviewed_at", "reviewed_by_id", "validity_due_date", "assigned_pm_id", "review_notes"):
            self.assertEqual(after[field], before[field], field)

    @patch("hub.views._send_sqr_for_revision_email")
    def test_manual_resume_choices_preserve_hold_reason_and_remarks(self, revision_email):
        for status in ("submitted", "for_revision", "reviewed"):
            with self.subTest(status=status):
                sqr = self.create_sqr()
                sqr.remarks = "Independent filing note"
                sqr.save()
                self.assertEqual(self.update_status("on_hold", submission=sqr, hold_reason="Customer delay").status_code, 200)
                response = self.update_status(status, submission=sqr)
                self.assertEqual(response.status_code, 200, response.content)
                self.assertEqual(response.json()["hold_reason"], "Customer delay")
                sqr.refresh_from_db()
                self.assertEqual(sqr.status, status)
                self.assertEqual(sqr.hold_reason, "Customer delay")
                self.assertEqual(sqr.remarks, "Independent filing note")
                if status == "reviewed":
                    self.assertIsNotNone(sqr.reviewed_at)
                    self.assertIsNotNone(sqr.validity_due_date)

    def test_hold_change_and_history_capture_reason_and_undo_restores_it(self):
        SqrSubmission.objects.filter(pk=self.submission.pk).update(hold_reason="Earlier retained reason")
        before = self.snapshot()
        response = self.update_status("on_hold", hold_reason="New hold reason")
        self.assertEqual(response.status_code, 200)
        change = SqrSubmissionChange.objects.get(pk=response.json()["change_id"])
        history = SqrSubmissionHistory.objects.get(submission=self.submission, action=SqrSubmissionHistory.Action.UPDATED)
        for entry in (change, history):
            self.assertEqual(entry.old_values["hold_reason"], "Earlier retained reason")
            self.assertEqual(entry.new_values["hold_reason"], "New hold reason")
            self.assertEqual(entry.old_values["status"], "submitted")
            self.assertEqual(entry.new_values["status"], "on_hold")
        undo = self.client.post(response.json()["undo_url"])
        self.assertEqual(undo.status_code, 200, undo.content)
        self.assertTrue(undo.json()["undone"])
        self.assertEqual(undo.json()["hold_reason"], "Earlier retained reason")
        after = self.snapshot()
        for field in ("status", "hold_reason", "remarks", "reviewed_at", "validity_due_date", "review_notes"):
            self.assertEqual(after[field], before[field], field)

    def test_history_restore_supports_hold_reason(self):
        self.assertEqual(self.update_status("on_hold", hold_reason="Pending decision").status_code, 200)
        history = SqrSubmissionHistory.objects.get(submission=self.submission, action=SqrSubmissionHistory.Action.UPDATED)
        response = self.client.post(reverse("hub:sqr-history-restore", args=[self.submission.pk, history.pk]))
        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.json()["ok"])
        self.submission.refresh_from_db()
        history.refresh_from_db()
        self.assertEqual(self.submission.status, "submitted")
        self.assertEqual(self.submission.hold_reason, "")
        self.assertEqual(self.submission.remarks, "SSE filing remarks <independent>")
        self.assertIsNotNone(history.restored_at)

    def test_stale_hold_reason_blocks_undo_and_history_restore(self):
        response = self.update_status("on_hold", hold_reason="Original delay")
        self.assertEqual(response.status_code, 200)
        history = SqrSubmissionHistory.objects.get(submission=self.submission, action=SqrSubmissionHistory.Action.UPDATED)
        SqrSubmission.objects.filter(pk=self.submission.pk).update(hold_reason="Later reason")
        before, counts = self.snapshot(), self.mutation_counts()
        for url in (response.json()["undo_url"], reverse("hub:sqr-history-restore", args=[self.submission.pk, history.pk])):
            with self.subTest(url=url):
                self.assert_rejected_without_mutation(self.client.post(url), before, counts, status=409)

    @patch("hub.views._send_sqr_new_submission_email")
    def test_filing_and_edit_remarks_are_optional_and_independent_for_siblings(self, send_email):
        self.client.force_login(self.engineer)
        for extra in ({}, {"remarks": ""}, {"remarks": "Only this submission has remarks"}):
            response = self.client.post(reverse("hub:sqr"), {**self.payload, **extra})
            self.assertEqual(response.status_code, 302, response.content)
        siblings = list(SqrSubmission.objects.filter(linked_request=self.rq).exclude(pk=self.submission.pk).order_by("pk"))
        self.assertEqual([s.remarks for s in siblings], ["", "", "Only this submission has remarks"])
        first, second, third = siblings
        SqrSubmission.objects.filter(pk=first.pk).update(status="on_hold", hold_reason="First only")
        response = self.client.post(reverse("hub:sqr-edit", args=[first.pk]), {**self.payload, "remarks": "Edited SSE remark"})
        self.assertEqual(response.status_code, 302, response.content)
        first.refresh_from_db()
        self.assertEqual(first.remarks, "Edited SSE remark")
        self.assertEqual(first.hold_reason, "First only")
        self.assertEqual(first.status, "on_hold")
        response = self.client.post(reverse("hub:sqr-edit", args=[first.pk]), self.payload)
        self.assertEqual(response.status_code, 302)
        first.refresh_from_db()
        self.assertEqual(first.remarks, "")
        self.assertEqual(first.hold_reason, "First only")
        second.refresh_from_db()
        third.refresh_from_db()
        self.assertEqual(second.remarks, "")
        self.assertEqual(second.status, "submitted")
        self.assertEqual(second.hold_reason, "")
        self.assertEqual(third.remarks, "Only this submission has remarks")
        self.assertEqual(third.hold_reason, "")

    def test_tracker_remarks_precede_status_and_hold_reason_is_in_status_cell(self):
        SqrSubmission.objects.filter(pk=self.submission.pk).update(status="on_hold", hold_reason="Customer <decision> pending")
        response = self.client.get(reverse("hub:sqr"))
        self.assertEqual(response.status_code, 200)
        parser = TrackerRows()
        parser.feed(response.content.decode())
        header = next(row for row in parser.rows if any(cell["text"].strip() == "SQR Status" for cell in row["cells"]))
        labels = [cell["text"].strip() for cell in header["cells"]]
        status_index = labels.index("SQR Status")
        self.assertEqual(labels[status_index - 1], "Remarks")
        row = next(row for row in parser.rows if row["attrs"].get("data-sqr-ref") == self.submission.reference_code)
        self.assertEqual(row["attrs"]["data-sqr-status"], "on_hold")
        self.assertIn(self.submission.remarks, row["cells"][status_index - 1]["text"])
        self.assertIn("On Hold", row["cells"][status_index]["text"])
        self.assertIn("Customer <decision> pending", row["cells"][status_index]["text"])
        self.assertContains(response, "Customer &lt;decision&gt; pending")
        content = response.content.decode()
        filter_markup = content.split('id="sqr-table-filter-status"', 1)[1].split("</select>", 1)[0]
        self.assertIn('<option value="on_hold">On Hold</option>', filter_markup)

    def export_workbook(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("hub:sqr-export"))
        self.assertEqual(response.status_code, 200)
        return load_workbook(io.BytesIO(response.content))

    def import_workbook(self, workbook, job_id):
        content = io.BytesIO()
        workbook.save(content)
        importer = SqrImportView()
        with patch("hub.views.close_old_connections"):
            importer._run_import_job(job_id=job_id, user_id=self.admin.pk, file_bytes=content.getvalue())
        return importer._get_status(self.admin, job_id)

    def test_export_import_roundtrip_preserves_status_reason_and_sibling_remarks(self):
        second = self.create_sqr()
        SqrSubmission.objects.filter(pk=self.submission.pk).update(status="on_hold", hold_reason="Customer delay")
        SqrSubmission.objects.filter(pk=second.pk).update(remarks="Second note", hold_reason="Earlier delay")
        expected = {
            self.submission.reference_code: ("on_hold", "Customer delay", self.submission.remarks),
            second.reference_code: ("submitted", "Earlier delay", "Second note"),
        }
        workbook = self.export_workbook()
        rows = list(workbook.active.iter_rows(values_only=True))
        headers = list(rows[0])
        self.assertEqual(headers[headers.index("SQR Status") - 1], "Remarks")
        self.assertIn("Hold Reason", headers)
        held_row = next(row for row in rows[1:] if row[headers.index("SQR ID")] == self.submission.reference_code)
        self.assertEqual(held_row[headers.index("SQR Status")], "On Hold")
        self.assertEqual(held_row[headers.index("Hold Reason")], "Customer delay")
        self.assertEqual(held_row[headers.index("Remarks")], self.submission.remarks)
        status = self.import_workbook(workbook, "hold-roundtrip-test")
        self.assertEqual(status["state"], "completed", status)
        imported = {
            sqr.reference_code: (sqr.status, sqr.hold_reason, sqr.remarks)
            for sqr in SqrSubmission.objects.filter(linked_request=self.rq)
        }
        self.assertEqual(imported, expected)

    def test_import_invalid_hold_reason_does_not_replace_existing_submissions(self):
        workbook = self.export_workbook()
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        sheet.cell(2, headers.index("SQR Status") + 1, "On Hold")
        for reason in ("", " \t "):
            with self.subTest(reason=reason):
                sheet.cell(2, headers.index("Hold Reason") + 1, reason)
                before, counts = self.snapshot(), self.mutation_counts()
                status = self.import_workbook(workbook, "hold-invalid-import-test")
                self.assertEqual(status["state"], "failed", status)
                self.assertEqual(self.snapshot(), before)
                self.assertEqual(self.mutation_counts(), counts)
                self.assertEqual(SqrSubmission.objects.count(), 1)

    def legacy_workbook(self):
        workbook = self.export_workbook()
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        for header in ("Hold Reason", "Remarks"):
            sheet.delete_cols(headers.index(header) + 1)
            headers.remove(header)
        self.assertEqual(len(headers), 43)
        self.assertEqual(headers[19], "SQR Status")
        self.assertEqual(headers[37], "PO Remarks")
        return workbook

    def test_named_legacy_headers_keep_status_and_po_remarks_in_original_columns(self):
        SqrSubmission.objects.filter(pk=self.submission.pk).update(
            status="for_revision", remarks="Not present in legacy file",
            revenue_remarks="Legacy PO note", revenue_overview="Legacy billing note",
        )
        workbook = self.legacy_workbook()
        status = self.import_workbook(workbook, "legacy-named-hold-test")
        self.assertEqual(status["state"], "completed", status)
        imported = SqrSubmission.objects.get(reference_code=self.submission.reference_code)
        self.assertEqual(imported.status, "for_revision")
        self.assertEqual(imported.hold_reason, "")
        self.assertEqual(imported.remarks, "")
        self.assertEqual(imported.revenue_remarks, "Legacy PO note")
        self.assertEqual(imported.revenue_overview, "Legacy billing note")

    def test_headerless_legacy_43_and_shorter_rows_use_legacy_positions(self):
        SqrSubmission.objects.filter(pk=self.submission.pk).update(
            status="for_revision", revenue_remarks="Headerless PO note",
        )
        for width in (43, 40):
            with self.subTest(width=width):
                workbook = self.legacy_workbook()
                sheet = workbook.active
                sheet.delete_rows(1)
                if width < 43:
                    sheet.delete_cols(width + 1, 43 - width)
                status = self.import_workbook(workbook, f"legacy-headerless-{width}-test")
                self.assertEqual(status["state"], "completed", status)
                imported = SqrSubmission.objects.get(reference_code=self.submission.reference_code)
                self.assertEqual(imported.status, "for_revision")
                self.assertEqual(imported.remarks, "")
                self.assertEqual(imported.hold_reason, "")
                self.assertEqual(imported.revenue_remarks, "Headerless PO note")

    def test_named_new_headers_preserve_three_independent_notes_fields(self):
        SqrSubmission.objects.filter(pk=self.submission.pk).update(
            status="on_hold", hold_reason="Pending scope", remarks="SSE note",
            revenue_remarks="Separate PO note",
        )
        original = self.export_workbook()
        headers = [cell.value for cell in original.active[1]]
        self.assertEqual(headers[headers.index("SQR Status") - 1], "Remarks")
        self.assertEqual(headers[headers.index("SQR Status") + 1], "Hold Reason")
        status = self.import_workbook(original, "named-new-hold-headers-test")
        self.assertEqual(status["state"], "completed", status)
        imported = SqrSubmission.objects.get(reference_code=self.submission.reference_code)
        self.assertEqual(imported.status, "on_hold")
        self.assertEqual(imported.hold_reason, "Pending scope")
        self.assertEqual(imported.remarks, "SSE note")
        self.assertEqual(imported.revenue_remarks, "Separate PO note")

    def test_stale_engineer_remarks_edit_preserves_concurrent_pm_hold_and_history(self):
        request = RequestFactory().post(
            reverse("hub:sqr-edit", args=[self.submission.pk]),
            {**self.payload, "remarks": "Updated engineer remarks"},
        )
        request.user = self.engineer
        request.session = self.client.session
        request._messages = FallbackStorage(request)
        view = SqrEngineerUpdateView()
        view.setup(request, pk=self.submission.pk)
        view.object = SqrSubmission.objects.get(pk=self.submission.pk)
        form = view.get_form()
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.instance.status, "submitted")
        self.assertEqual(form.instance.hold_reason, "")

        response = self.update_status("on_hold", hold_reason="PM hold after engineer loaded form")
        self.assertEqual(response.status_code, 200)
        hold_change = SqrSubmissionChange.objects.get(pk=response.json()["change_id"])
        response = view.form_valid(form)
        self.assertEqual(response.status_code, 302)
        self.submission.refresh_from_db()
        self.assertEqual(self.submission.remarks, "Updated engineer remarks")
        self.assertEqual(self.submission.status, "on_hold")
        self.assertEqual(self.submission.hold_reason, "PM hold after engineer loaded form")
        hold_change.refresh_from_db()
        self.assertEqual(hold_change.new_values["status"], "on_hold")
        self.assertEqual(hold_change.new_values["hold_reason"], self.submission.hold_reason)
        history = SqrSubmissionHistory.objects.filter(submission=self.submission).latest("pk")
        self.assertEqual(history.new_values["remarks"], "Updated engineer remarks")
        for field in ("status", "hold_reason"):
            self.assertNotIn(field, history.old_values)
            self.assertNotIn(field, history.new_values)

    def update_hold_reason(self, value):
        return self.client.post(
            reverse("hub:sqr-inline-update", args=[self.submission.pk]),
            data=json.dumps({"field": "hold_reason", "value": value}),
            content_type="application/json",
        )

    def test_direct_hold_reason_invalid_values_are_rejected_atomically(self):
        for user in (self.pm, self.admin):
            self.client.force_login(user)
            for status in ("submitted", "on_hold"):
                SqrSubmission.objects.filter(pk=self.submission.pk).update(
                    status=status, hold_reason="Retained reason",
                )
                invalid_values = [None, False, 0, [], {}]
                if status == "on_hold":
                    invalid_values.extend(["", " \t\r\n "])
                for value in invalid_values:
                    with self.subTest(role=user.role, status=status, value=value):
                        before, counts = self.snapshot(), self.mutation_counts()
                        self.assert_rejected_without_mutation(
                            self.update_hold_reason(value), before, counts,
                        )

    def test_pm_admin_direct_hold_reason_is_trimmed_without_changing_status_or_remarks(self):
        for user in (self.pm, self.admin):
            self.client.force_login(user)
            for status in ("submitted", "on_hold"):
                with self.subTest(role=user.role, status=status):
                    SqrSubmission.objects.filter(pk=self.submission.pk).update(
                        status=status, hold_reason="Previous reason",
                    )
                    before = self.snapshot()
                    response = self.update_hold_reason(" \t Customer confirmation pending \r\n ")
                    self.assertEqual(response.status_code, 200, response.content)
                    self.assertTrue(response.json()["ok"])
                    self.assertEqual(response.json()["hold_reason"], "Customer confirmation pending")
                    after = self.snapshot()
                    self.assertEqual(after["hold_reason"], "Customer confirmation pending")
                    for field in ("status", "remarks", "review_notes", "reviewed_at", "validity_due_date"):
                        self.assertEqual(after[field], before[field], field)
                    change = SqrSubmissionChange.objects.get(pk=response.json()["change_id"])
                    self.assertEqual(change.field, "hold_reason")
                    self.assertEqual(change.old_values["hold_reason"], "Previous reason")
                    self.assertEqual(change.new_values["hold_reason"], "Customer confirmation pending")

    def test_sse_cannot_directly_edit_hold_reason(self):
        SqrSubmission.objects.filter(pk=self.submission.pk).update(
            status="on_hold", hold_reason="PM reason",
        )
        for user in (self.engineer, self.paused_engineer):
            with self.subTest(role=user.role):
                self.client.force_login(user)
                before, counts = self.snapshot(), self.mutation_counts()
                self.assert_rejected_without_mutation(
                    self.update_hold_reason("Unauthorized reason"), before, counts, status=403,
                )

    def test_new_remarks_never_fall_back_into_empty_po_remarks(self):
        SqrSubmission.objects.filter(pk=self.submission.pk).update(
            remarks="SSE note only", revenue_remarks="",
        )
        status = self.import_workbook(self.export_workbook(), "independent-import-remarks-test")
        self.assertEqual(status["state"], "completed", status)
        imported = SqrSubmission.objects.get(reference_code=self.submission.reference_code)
        self.assertEqual(imported.remarks, "SSE note only")
        self.assertEqual(imported.revenue_remarks, "")


class SqrOnHoldMigrationTests(MultipleSqrFixtures, TransactionTestCase):
    def test_0053_preserves_existing_status_remarks_and_approval_metadata(self):
        first = self.create_sqr()
        reviewed_at = timezone.now() - timedelta(days=3)
        due_date = timezone.localdate() + timedelta(days=27)
        SqrSubmission.objects.filter(pk=first.pk).update(
            status="reviewed", remarks="Existing SSE remarks", revenue_remarks="Existing PO remarks",
            reviewed_at=reviewed_at, reviewed_by=self.pm, validity_due_date=due_date,
        )
        executor = MigrationExecutor(connection)
        leaves = executor.loader.graph.leaf_nodes()
        current = [node for node in executor.loader.graph.nodes if node[0] == "hub" and node[1].startswith("0053_")]
        self.assertEqual(len(current), 1)
        previous = [("hub", "0052_allow_multiple_sqrs_per_request")]
        try:
            executor.migrate(previous)
            executor = MigrationExecutor(connection)
            old_sqr = executor.loader.project_state(previous).apps.get_model("hub", "SqrSubmission")
            before = old_sqr.objects.filter(pk=first.pk).values().get()
            executor.migrate(current)
            historical_sqr = executor.loader.project_state(current).apps.get_model("hub", "SqrSubmission")
            after = historical_sqr.objects.filter(pk=first.pk).values().get()
            for field, value in before.items():
                self.assertEqual(after[field], value, field)
            self.assertEqual(after["hold_reason"], "")
            self.assertTrue(historical_sqr._meta.get_field("hold_reason").blank)
            self.assertIn(("on_hold", "On Hold"), historical_sqr._meta.get_field("status").choices)
        finally:
            MigrationExecutor(connection).migrate(leaves)
