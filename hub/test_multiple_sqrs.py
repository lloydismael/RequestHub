import io
import json
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from threading import Barrier
from unittest.mock import patch

from django.db import IntegrityError, close_old_connections, connection
from django.db.migrations.executor import MigrationExecutor
from django.test import RequestFactory, TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from hub.models import Account, Notification, Request, SqrProjectPerformance, SqrSubmission
from hub.views import RequestCollaborativeManageView, SqrEngineerUpdateView, SqrImportView, _make_sqr_edit_form


class MultipleSqrFixtures:
    def setUp(self):
        super().setUp()
        self.engineer = User.objects.create_user(username="multi_sqr_engineer", role=User.Roles.ENGINEER)
        self.other = User.objects.create_user(username="multi_sqr_other", role=User.Roles.ENGINEER)
        self.pm = User.objects.create_user(username="multi_sqr_pm", role=User.Roles.PM_ESG)
        self.requestor = User.objects.create_user(username="multi_sqr_requestor", role=User.Roles.REQUESTOR)
        self.admin = User.objects.create_user(username="multi_sqr_admin", role=User.Roles.ADMIN)
        self.account = Account.objects.create(name="Multiple SQR Account")
        self.rq = Request.objects.create(
            requestor=self.requestor, account=self.account, account_manager="Contact",
            product_category="Azure", engagement_type=Request.Engagement.SUPPORT,
            priority=Request.Priority.MEDIUM, engineer=self.engineer,
        )
        self.payload = {
            "linked_request": str(self.rq.pk), "customer_name": self.account.name,
            "customer_company": "ESS", "customer_contact": "",
            "pm_esg_reviewer": str(self.pm.pk), "project_title": "Independent project",
            "project_details": "Deployment Only", "sse_manhrs": "20",
            "sqr_folder_link": "https://example.com/quotation",
        }

    def create_sqr(self, rq=None, engineer=None):
        return SqrSubmission.objects.create(
            linked_request=rq or self.rq, engineer=engineer or self.engineer,
            pm_esg_reviewer=self.pm, customer_name=self.account.name,
            customer_company="ESS", project_title="Independent project",
            project_details="Implementation", sse_manhrs=Decimal("20"),
            sqr_folder_link="https://example.com/quotation",
        )


@override_settings(ACS_EMAIL_CONNECTION_STRING="", ACS_EMAIL_SENDER="")
class MultipleSqrTests(MultipleSqrFixtures, TestCase):
    def test_model_allows_shared_request_and_independent_lifecycles(self):
        first = self.create_sqr()
        second = self.create_sqr()
        self.assertNotEqual(first.reference_code, second.reference_code)
        self.assertEqual(second.sequence_number, first.sequence_number + 1)
        first.status = SqrSubmission.Status.APPROVED
        first.save()
        second.refresh_from_db()
        self.assertEqual(second.status, SqrSubmission.Status.FOR_PROCESSING)
        SqrProjectPerformance.objects.create(submission=first)
        SqrProjectPerformance.objects.create(submission=second)
        self.assertEqual(SqrProjectPerformance.objects.filter(submission__linked_request=self.rq).count(), 2)
        first.delete()
        self.assertTrue(SqrSubmission.objects.filter(pk=second.pk).exists())
        self.assertTrue(Request.objects.filter(pk=self.rq.pk).exists())

    @patch("hub.views._send_sqr_new_submission_email")
    def test_repeated_posts_keep_request_available_and_notify_per_sqr(self, send_email):
        self.client.force_login(self.engineer)
        for _ in range(3):
            response = self.client.post(reverse("hub:sqr"), self.payload)
            self.assertEqual(response.status_code, 302)
            response = self.client.get(reverse("hub:sqr"))
            self.assertIn(self.rq, response.context["sqr_form"].fields["linked_request"].queryset)
        sqrs = list(SqrSubmission.objects.filter(linked_request=self.rq))
        self.assertEqual(len(sqrs), 3)
        self.assertEqual(len({s.reference_code for s in sqrs}), 3)
        self.assertEqual(send_email.call_count, 3)
        self.assertEqual(Notification.objects.filter(recipient=self.pm, source="SQR · New Submission").count(), 3)

    def test_invalid_post_retains_reusable_request(self):
        self.create_sqr()
        self.client.force_login(self.engineer)
        response = self.client.post(reverse("hub:sqr"), {**self.payload, "project_title": ""})
        self.assertEqual(response.status_code, 200)
        self.assertIn("project_title", response.context["sqr_form"].errors)
        self.assertIn(self.rq, response.context["sqr_form"].fields["linked_request"].queryset)
        self.assertEqual(SqrSubmission.objects.filter(linked_request=self.rq).count(), 1)

    def test_other_and_backup_engineer_cannot_create_for_request(self):
        self.rq.backup_engineer = self.other
        self.rq.save()
        self.client.force_login(self.other)
        response = self.client.post(reverse("hub:sqr"), self.payload)
        self.assertEqual(response.status_code, 200)
        self.assertIn("linked_request", response.context["sqr_form"].errors)
        self.assertFalse(SqrSubmission.objects.filter(linked_request=self.rq).exists())
        response = self.client.get(reverse("hub:sqr"))
        self.assertNotIn(self.rq, response.context["sqr_form"].fields["linked_request"].queryset)

    def test_non_engineer_creation_permissions_unchanged(self):
        for user in (self.pm, self.admin, self.requestor):
            self.client.force_login(user)
            self.assertEqual(self.client.post(reverse("hub:sqr"), self.payload).status_code, 302)
        self.assertFalse(SqrSubmission.objects.filter(linked_request=self.rq).exists())

    def test_siblings_edit_independently_and_refresh_account_map(self):
        first, second = self.create_sqr(), self.create_sqr()
        self.client.force_login(self.engineer)
        response = self.client.post(reverse("hub:sqr-edit", args=[first.pk]), {
            **self.payload, "project_title": "Changed only first",
        })
        self.assertEqual(response.status_code, 302)
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual(first.project_title, "Changed only first")
        self.assertEqual(second.project_title, "Independent project")
        self.assertEqual(first.history_entries.count(), 1)
        self.assertEqual(second.history_entries.count(), 0)
        self.assertIn(self.rq, _make_sqr_edit_form(self.engineer).fields["linked_request"].queryset)
        request = RequestFactory().get(reverse("hub:sqr-edit", args=[first.pk]))
        request.user = self.admin
        view = SqrEngineerUpdateView()
        view.setup(request, pk=first.pk)
        view.object = first
        form = view.get_form()
        self.assertIn(self.rq, form.fields["linked_request"].queryset)
        mapping = json.loads(form.fields["linked_request"].widget.attrs["data-account-map"])
        self.assertEqual(mapping[str(self.rq.pk)], self.account.name)

    def test_reassigned_historical_link_can_still_be_edited_not_created(self):
        first, second = self.create_sqr(), self.create_sqr()
        self.rq.engineer = self.other
        self.rq.save()
        self.client.force_login(self.engineer)
        self.assertEqual(self.client.post(reverse("hub:sqr-edit", args=[first.pk]), self.payload).status_code, 302)
        self.assertIn(self.rq, _make_sqr_edit_form(self.engineer).fields["linked_request"].queryset)
        response = self.client.post(reverse("hub:sqr"), self.payload)
        self.assertIn("linked_request", response.context["sqr_form"].errors)
        self.assertEqual(SqrSubmission.objects.filter(linked_request=self.rq).count(), 2)

    def test_request_management_shows_all_authorized_siblings(self):
        first, second = self.create_sqr(), self.create_sqr()
        hidden = self.create_sqr(engineer=self.other)
        self.client.force_login(self.engineer)
        response = self.client.get(reverse("hub:request-manage-collab", args=[self.rq.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(list(response.context["linked_sqrs"]), [second, first])
        self.assertContains(response, f"View linked SQR {first.reference_code}")
        self.assertContains(response, f"View linked SQR {second.reference_code}")
        self.assertNotContains(response, hidden.reference_code)
        request = RequestFactory().get("/")
        request.user = self.admin
        view = RequestCollaborativeManageView()
        view.setup(request)
        self.assertEqual(list(view.get_context_data(self.rq)["linked_sqrs"]), [hidden, second, first])
        request.user = self.requestor
        self.assertFalse(view.get_context_data(self.rq)["linked_sqrs"].exists())

    def test_export_and_import_preserve_sibling_rows(self):
        first, second = self.create_sqr(), self.create_sqr()
        self.client.force_login(self.admin)
        response = self.client.get(reverse("hub:sqr-export"))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        rows = list(workbook.active.iter_rows(values_only=True))
        headers = list(rows[0])
        sqr_col, rq_col = headers.index("SQR ID"), headers.index("RQ ID")
        siblings = [row for row in rows[1:] if row[rq_col] == self.rq.reference_code]
        self.assertEqual({row[sqr_col] for row in siblings}, {first.reference_code, second.reference_code})
        importer = SqrImportView()
        with patch("hub.views.close_old_connections"):
            importer._run_import_job(
                job_id="multi-sqr-import-test", user_id=self.admin.pk, file_bytes=response.content,
            )
        status = importer._get_status(self.admin, "multi-sqr-import-test")
        self.assertEqual(status["state"], "completed", status)
        self.assertEqual(SqrSubmission.objects.filter(linked_request=self.rq).count(), 2)

    def test_request_management_zero_and_single_link(self):
        self.client.force_login(self.engineer)
        url = reverse("hub:request-manage-collab", args=[self.rq.pk])
        response = self.client.get(url)
        self.assertEqual(list(response.context["linked_sqrs"]), [])
        first = self.create_sqr()
        response = self.client.get(url)
        self.assertEqual(list(response.context["linked_sqrs"]), [first])
        self.assertContains(response, f"View linked SQR {first.reference_code}")

    def test_database_constraint_removed(self):
        with connection.cursor() as cursor:
            constraints = connection.introspection.get_constraints(cursor, SqrSubmission._meta.db_table)
        self.assertNotIn("unique_sqrsubmission_linked_request", constraints)

    @patch("hub.views._send_sqr_new_submission_email")
    @patch("hub.views.logger.exception")
    @patch("hub.views.save_configured_instance", side_effect=IntegrityError("save failure"))
    def test_integrity_failure_is_not_reported_as_duplicate_request(self, save, log, email):
        self.client.force_login(self.engineer)
        response = self.client.post(reverse("hub:sqr"), self.payload)
        self.assertEqual(response.status_code, 200)
        self.assertIn("The SQR could not be saved. Please try again.", response.context["sqr_form"].non_field_errors())
        self.assertContains(response, "The SQR could not be saved. Please try again.")
        self.assertFalse(SqrSubmission.objects.filter(linked_request=self.rq).exists())
        email.assert_not_called()


class ConcurrentMultipleSqrTests(MultipleSqrFixtures, TransactionTestCase):
    def test_migration_preserves_existing_submission(self):
        previous = [("hub", "0051_allow_admin_requestor")]
        current = [("hub", "0052_allow_multiple_sqrs_per_request")]
        first = self.create_sqr()
        original = (first.pk, first.reference_code, first.linked_request_id, first.project_title)
        executor = MigrationExecutor(connection)
        executor.migrate(previous)
        try:
            executor = MigrationExecutor(connection)
            executor.migrate(current)
            first.refresh_from_db()
            self.assertEqual((first.pk, first.reference_code, first.linked_request_id, first.project_title), original)
            second = self.create_sqr()
            self.assertNotEqual(first.reference_code, second.reference_code)
        finally:
            MigrationExecutor(connection).migrate(current)

    def test_concurrent_numbering_same_and_different_requests(self):
        if connection.vendor != "postgresql":
            self.skipTest("Production PostgreSQL advisory-lock assertion")
        second_rq = Request.objects.create(
            requestor=self.requestor, account=self.account, account_manager="Contact",
            product_category="Azure", engagement_type=Request.Engagement.SUPPORT,
            priority=Request.Priority.MEDIUM, engineer=self.engineer,
        )
        for requests in ((self.rq, self.rq), (self.rq, second_rq)):
            barrier = Barrier(2)

            def create(rq):
                close_old_connections()
                try:
                    barrier.wait(timeout=10)
                    return self.create_sqr(rq=rq).reference_code
                finally:
                    close_old_connections()

            with ThreadPoolExecutor(max_workers=2) as executor:
                codes = list(executor.map(create, requests))
            self.assertEqual(len(set(codes)), 2)
        self.assertEqual(SqrSubmission.objects.count(), 4)
